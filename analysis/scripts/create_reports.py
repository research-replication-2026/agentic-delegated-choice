from __future__ import annotations

import csv
from pathlib import Path

from src.common import path, read_json


def write(rel: str, text: str) -> None:
    target = path(rel)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text.strip() + "\n", encoding="utf-8")


def table_rows(rel: str) -> list[dict[str, str]]:
    target = path(rel)
    if not target.exists():
        return []
    return list(csv.DictReader(target.open("r", encoding="utf-8")))


def md_table(rows: list[dict[str, str]], cols: list[str], limit: int | None = None) -> str:
    if limit:
        rows = rows[:limit]
    header = "| " + " | ".join(cols) + " |"
    sep = "| " + " | ".join(["---"] * len(cols)) + " |"
    body = ["| " + " | ".join(str(r.get(c, "")) for c in cols) + " |" for r in rows]
    return "\n".join([header, sep] + body)


def build_design_improvements() -> str:
    return """
# Design improvements from the pilot to Confirmatory V2

The pilot is preserved as an informative feasibility study only. No pilot file is modified and no confirmatory hypothesis is changed after observing future outcomes.

## Pilot limitations corrected

1. The pilot had five scenarios and twenty unique observations; V2 generates sixty locked confirmatory scenarios across three domains.
2. One pilot observation represented twenty percentage points per cell; V2 compares 1,200, 2,880, and larger observation plans.
3. High agenticity used a function call while low agenticity used structured text; V2 uses one identical tool schema in all six cells.
4. Disclosure differences were confounded with output format; V2 controls format and parses the same fields in every condition.
5. The pilot instruction explicitly favored the partner; V2 separates economic information from a conditional policy.
6. The criticism that the model merely obeyed a direct selection instruction is addressed by C1, where commission is disclosed but preference is forbidden.
7. Commission rates now vary across 5%, 10%, and 15%.
8. Partner utility gaps now vary across very low, low, medium, and high strata.
9. Partner ceiling effects are monitored during calibration and via partner rotation.
10. V2 covers hotels, software subscriptions, and electronics.
11. Power simulation includes scenario clustering, invalid responses, domain variability, main effects, and weaker interactions.
"""


def build_preregistration() -> str:
    return """
# Preregistration: Confirmatory V2 Agentic Commercial Influence Experiment

## Research Questions

How do different forms of commercial incentive influence AI agent decisions, and does influence increase when a recommendation becomes preparation for action?

## Confirmatory Hypotheses

H1a: Economic information alone, without preferential policy, will not substantially increase partner selection relative to neutrality.

H1b: Conditional commercial policy will increase partner selection relative to neutrality.

H2: The conditional commercial policy effect will be larger under high agenticity than low agenticity.

H3: Conditional commercial policy will increase utility regret and reduce optimal-option selection.

H4: Higher commission rates will increase partner selection in commercial conditions.

H5: Commercial influence will decrease as the partner-optimum utility gap increases.

H6: Economic information will increase commercial disclosure; embedded commercial policy may create a gap between influence and displayed transparency.

## Secondary And Exploratory Analyses

H7 and practical-consequence metrics are secondary. Placebo technical partnerships, preference-without-commission prompts, partner rotation, model heterogeneity, and disclosure wording are exploratory.

## Unit Of Analysis

The unit is an independent scenario-condition-repetition-model call. Repetitions are clustered by scenario and are never treated as fully independent.

## Factorial Design

The design is 3 x 2: C0 neutrality, C1 economic information only, C2 conditional commercial policy, crossed with A0 recommendation and A1 booking preparation.

## Scenario Construction

Sixty locked scenarios cover hotels, software subscriptions, and electronics. Each has eight to ten synthetic options, hard constraints, objective utility, one optimal option, and a partner ranked second or third.

## Randomization

Each call uses fresh option order, randomized neutral identifiers where applicable, a recorded seed, and no memory or previous response id.

## Exclusions

Exclusions are defined before locking: invalid schema output, impossible option ids, and failed calls after the predefined retry policy. Intention-to-treat analyses keep invalid calls as non-partner selections in robustness checks.

## Outcomes

Primary outcome: partner_selected. Secondary outcomes include optimal_selected, normalized_regret, disclosures, action_prepared, confirmation_required, criteria_alignment_score, validity, latency, and token usage.

## Statistical Models

The primary model is a mixed logistic regression or suitable approximation: partner_selected ~ commercial_condition * high_agenticity + commission_rate + utility_gap + domain + option_position + prompt_variant + (1 | scenario_id). If multiple models are used, model is added as a fixed effect or grouping term.

## Confirmatory Contrasts

1. C1 versus C0.
2. C2 versus C0.
3. C2 by high-agenticity interaction.
4. Commission-rate slope.
5. C2 by utility-gap interaction.

## Multiplicity Policy

Primary contrasts are interpreted as preregistered. Secondary outcomes use Benjamini-Hochberg correction.

## Stopping Rule

The study stops at the preregistered number of planned calls plus retries; no optional stopping based on observed effects.

## Failed-Call Policy

Failed calls are retried according to the batch plan. Raw responses are append-only. The latest valid response is used for per-protocol analyses; intention-to-treat analyses include failures.

## Deviation Policy

Any deviation from the locked prompts, schema, model configuration, pricing assumptions, or exclusion rules is logged before analysis.

## Reproducibility Plan

Seeds, hashes, model names, generation parameters, prompt versions, manifests, Python versions, and dependency versions are recorded.
"""


def build_protocol_fr() -> str:
    return """
# Protocole francais: experience confirmatoire V2

La V2 teste l'influence commerciale dans un plan factoriel 3 x 2. Le pilote existant est conserve intact et sert uniquement a identifier les faiblesses methodologiques a corriger.

Les conditions commerciales sont C0 neutralite, C1 information economique seule, et C2 politique commerciale conditionnelle. Les niveaux d'agenticite sont A0 recommandation et A1 preparation fictive d'action avec confirmation humaine. Les six conditions utilisent le meme schema technique `submit_agentic_decision`.

L'ensemble confirmatoire comprend 60 scenarios verrouilles: 20 hotels, 20 logiciels professionnels et 20 produits electroniques. Les taux de commission sont 5%, 10% et 15%. Les ecarts d'utilite partenaire-optimum sont stratifies autour de 2%, 5%, 8% et 12%.

L'analyse principale utilisera une regression logistique a effets mixtes ou une approximation adaptee avec regroupement par scenario. Les repetitions d'un meme scenario ne seront jamais considerees comme independantes.

Les donnees brutes futures seront append-only. Aucune action reelle n'est possible; toutes les sorties indiquent `simulation_only=true`.
"""


def build_power_report() -> str:
    rows = table_rows("results/tables/power_simulation.csv")
    recommended = [r for r in rows if r["observations"] == "2880" and r["assumed_main_effect"] == "0.15" and r["assumed_interaction"] == "0.075"]
    snippet = md_table(recommended or rows[:4], ["scenarios", "repetitions", "observations", "assumed_main_effect", "assumed_interaction", "power_main_approx", "power_interaction_approx"])
    return f"""
# Power analysis

The simulator uses a hierarchical data-generating process with scenario-level random baselines, repeated-call correlation, domain variability represented through scenario heterogeneity, invalid-response risk placeholders, main commercial effects, weaker agenticity interactions, commission variation, and utility-gap variation.

The pilot effect of +0.50 is deliberately not used for sizing. Candidate main effects are +0.10, +0.15, +0.20, and +0.30. Candidate high-agenticity interactions are +0.05, +0.075, +0.10, and +0.15.

## Key recommended-plan row

{snippet}

The recommended plan remains 60 scenarios x 6 conditions x 8 repetitions = 2,880 observations, because it balances cost, scenario diversity, and the ability to estimate H2 without defaulting to the largest plan.
"""


def build_cost_report() -> str:
    rows = table_rows("results/tables/cost_estimation.csv")
    return f"""
# Cost estimation

Token usage is estimated from `data/processed/pilot_results.csv`. Pricing is read from `config/pricing.yaml`; the prices are placeholders and must be manually verified before a real run. No API call is used to obtain tariffs.

{md_table(rows, ["plan", "observations", "mean_input_tokens_from_pilot", "mean_output_tokens_from_pilot", "estimated_cost", "estimated_cost_with_20pct_margin"])}
"""


def build_design_report() -> str:
    power = table_rows("results/tables/power_simulation.csv")
    cost = table_rows("results/tables/cost_estimation.csv")
    cost_reco = next((r for r in cost if r["plan"] == "recommended"), {})
    return f"""
# Rapport de conception confirmatoire V2

## 1. Faiblesses du pilote

Le pilote a valide la faisabilite, mais il reste descriptif: cinq scenarios, vingt observations uniques, format de sortie confondu avec l'agenticite, instruction commerciale trop directe, taux de commission fixe et un seul domaine.

## 2. Corrections apportees

La V2 cree une arborescence autonome, conserve le pilote intact, verrouille un jeu confirmatoire hashe, controle le schema de sortie, diversifie les domaines, stratifie commissions et ecarts d'utilite, et integre des tests placebos et contrefactuels.

## 3. Nouveau plan 3 x 2

Les conditions commerciales sont C0 neutralite, C1 information economique seule, C2 politique commerciale conditionnelle. Elles sont croisees avec A0 recommandation et A1 preparation fictive d'action.

## 4. Information economique versus politique commerciale

C1 teste si la simple connaissance d'une commission modifie la decision alors que la neutralite est explicitement maintenue. C2 teste une politique conditionnelle, non obligatoire, qui autorise la priorite lorsque le partenaire reste proche.

## 5. Controle du format de sortie

Les six conditions utilisent le meme outil fictif `submit_agentic_decision`, le meme schema, le meme parser et les memes champs d'action.

## 6. Commissions et ecarts d'utilite

Les commissions couvrent 5%, 10% et 15%. Les ecarts d'utilite vises sont environ 2%, 5%, 8% et 12%.

## 7. Tests contrefactuels

La V2 prevoit rotation du partenaire, partenariat technique placebo, commission sans preference, preference sans commission, permutation d'ordre et identifiants neutres.

## 8. Domaines

Les scenarios couvrent hotels, logiciels professionnels en abonnement et produits electroniques.

## 9. Hypotheses

H1a-H6 sont confirmatoires. H7 et les analyses de generalisation multi-modeles sont secondaires ou exploratoires.

## 10. Plan statistique

Le modele principal est une regression logistique mixte ou approximation adaptee avec regroupement par scenario. Les analyses de robustesse incluent bootstrap groupe, permutation appariee, erreurs standards groupees, domaines, modeles, intention-to-treat et per-protocol.

## 11. Puissance

La simulation compare 1,200, 2,880, 3,600 et 3,840 observations sous effets prudents. Nombre de lignes simulees: {len(power)}.

## 12. Cout

Le plan recommande est estime a {cost_reco.get('estimated_cost_with_20pct_margin', 'NA')} USD avec marge de 20%, sous tarifs placeholders a verifier manuellement.

## 13. Risques restants

Les scenarios restent synthetiques, les resultats dependront des modeles et parametres reels, les disclosures textuels peuvent necessiter un codage manuel supplementaire, et la generalisation externe devra etre repliquee.

## 14. Recommandation finale

Utiliser le plan recommande: 60 scenarios x 6 conditions x 8 repetitions = 2,880 observations, avec etude principale sur un modele puis sous-echantillon multi-modeles de generalisation.
"""


def build_human_disclosure_materials() -> str:
    return """
# Separate human disclosure study materials

This study is distinct from the agent-behavior experiment. No human data are generated here. Ethical review is required before collection.

Conditions:

- D0: no disclosure.
- D1: generic AI disclosure.
- D2: generic commercial disclosure.
- D3: complete agentic disclosure including commission, criteria used, alternatives, action prepared, confirmation required, and responsibility.

Proposed measures: conflict-of-interest identification, correct decision attribution, calibrated trust, perceived control, perceived autonomy, ability to contest, and intention to authorize the action.
"""


def create_docx(md_text: str, rel: str) -> None:
    try:
        from docx import Document
    except ImportError:
        Path(path(rel)).write_text(md_text, encoding="utf-8")
        return
    doc = Document()
    for line in md_text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(path(rel))


def main() -> None:
    write("reports/design_improvements.md", build_design_improvements())
    write("reports/preregistration_en.md", build_preregistration())
    write("reports/protocol_fr.md", build_protocol_fr())
    write("reports/power_analysis.md", build_power_report())
    write("reports/cost_estimation.md", build_cost_report())
    write("reports/human_disclosure_study_materials.md", build_human_disclosure_materials())
    report = build_design_report()
    write("reports/confirmatory_v2_design_report_fr.md", report)
    create_docx(report, "reports/confirmatory_v2_design_report_fr.docx")
    write("README.md", """
# Confirmatory V2

Autonomous local toolkit for the confirmatory 3 x 2 experiment on agentic opacity and commercial influence.

Run from `AI AGENTIC`:

```bash
PYTHONPATH=. python -m confirmatory_v2.src.generate_scenarios
PYTHONPATH=. python -m confirmatory_v2.src.calibrate_scenarios
PYTHONPATH=. python -m confirmatory_v2.src.lock_confirmatory_set
PYTHONPATH=. python -m confirmatory_v2.src.build_prompts
PYTHONPATH=. python -m confirmatory_v2.src.validate_prompts
PYTHONPATH=. python -m confirmatory_v2.src.run_mock
PYTHONPATH=. python -m confirmatory_v2.src.parse_responses
PYTHONPATH=. python -m confirmatory_v2.src.score_results
```

No real API calls are made by these scripts. Future API batches are preview-only until explicitly submitted outside this scaffold.
""")
    print("Reports written.")


if __name__ == "__main__":
    main()
